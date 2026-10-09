"""
Views for the Navratri Child Gift & Entry Management System.
"""
import csv
import io
import json
from datetime import datetime

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db import IntegrityError, transaction
from django.db.models import Q
from django.http import JsonResponse, HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.http import require_POST
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from .forms import ChildEditForm, ChildRegistrationForm, EventConfigForm
from .models import Child, DailyAttendance, DailyGift, EventConfig


# ─────────────────────────── Helper ──────────────────────────────────────────

def _get_config():
    return EventConfig.get_config()


# ─────────────────────────── Dashboard ───────────────────────────────────────

@login_required
def dashboard(request):
    config = _get_config()
    day = config.current_day

    total_children = Child.objects.count()
    today_entries = DailyAttendance.objects.filter(event_day=day).count()
    today_gifts = DailyGift.objects.filter(event_day=day, status='GIVEN').count()

    entered_ids = DailyAttendance.objects.filter(event_day=day).values_list('child_id', flat=True)
    gifted_ids = DailyGift.objects.filter(event_day=day, status='GIVEN').values_list('child_id', flat=True)

    not_entered = total_children - today_entries
    not_gifted = total_children - today_gifts

    recent_entries = (
        DailyAttendance.objects.filter(event_day=day)
        .select_related('child')
        .order_by('-entry_time')[:10]
    )
    recent_gifts = (
        DailyGift.objects.filter(event_day=day)
        .select_related('child')
        .order_by('-gift_time')[:10]
    )

    return render(request, 'events/dashboard.html', {
        'config': config,
        'day': day,
        'total_children': total_children,
        'today_entries': today_entries,
        'today_gifts': today_gifts,
        'not_entered': not_entered,
        'not_gifted': not_gifted,
        'recent_entries': recent_entries,
        'recent_gifts': recent_gifts,
        'entry_pct': int((today_entries / total_children * 100) if total_children else 0),
        'gift_pct': int((today_gifts / total_children * 100) if total_children else 0),
    })


# ─────────────────────────── Registration ────────────────────────────────────

@login_required
def register_child(request):
    if request.method == 'POST':
        form = ChildRegistrationForm(request.POST, request.FILES)
        if form.is_valid():
            child = form.save()
            messages.success(request, f'Registration successful! ID: {child.registration_id}')
            return redirect('registration_success', pk=child.pk)
    else:
        form = ChildRegistrationForm()

    return render(request, 'events/register.html', {'form': form})


@login_required
def registration_success(request, pk):
    child = get_object_or_404(Child, pk=pk)
    return render(request, 'events/registration_success.html', {'child': child})


@login_required
def id_card(request, pk):
    child = get_object_or_404(Child, pk=pk)
    return render(request, 'events/id_card.html', {'child': child})


@login_required
def edit_child(request, pk):
    child = get_object_or_404(Child, pk=pk)
    if request.method == 'POST':
        form = ChildEditForm(request.POST, request.FILES, instance=child)
        if form.is_valid():
            # If no new photo uploaded, keep existing
            if not request.FILES.get('photo') and child.photo:
                form.instance.photo = child.photo
            form.save()
            messages.success(request, f'{child.registration_id} updated successfully.')
            return redirect('child_profile', pk=child.pk)
    else:
        form = ChildEditForm(instance=child)

    return render(request, 'events/edit_child.html', {'form': form, 'child': child})


@login_required
def delete_child(request, pk):
    child = get_object_or_404(Child, pk=pk)
    if request.method == 'POST':
        name = child.name
        reg_id = child.registration_id
        child.delete()
        messages.success(request, f'{reg_id} ({name}) has been deleted.')
        return redirect('child_list')
    return render(request, 'events/delete_child.html', {'child': child})


@login_required
def child_list(request):
    children = Child.objects.all().order_by('registration_id')
    return render(request, 'events/child_list.html', {'children': children})


@login_required
def child_profile(request, pk):
    child = get_object_or_404(Child, pk=pk)
    history = child.get_history()
    config = _get_config()
    return render(request, 'events/child_profile.html', {
        'child': child,
        'history': history,
        'config': config,
    })


def _find_child_by_query(query_id):
    """
    Find child by raw search query or ID.
    Supports:
    - Exact match: e.g. '2026-1'
    - Plain numbers: e.g. '1', '2', '25' -> matches '2026-1', '2026-2', '2026-25' as well as legacy '1'
    - '2026-X': also matches legacy 'X'
    """
    if not query_id:
        return None
    raw = str(query_id).strip()
    if not raw:
        return None

    # Check candidates in priority order
    candidates = [raw]
    if raw.isdigit():
        candidates.append(f"2026-{raw}")
    elif raw.startswith("2026-"):
        candidates.append(raw[5:])

    for cand in candidates:
        child = Child.objects.filter(registration_id__iexact=cand).first()
        if child:
            return child

    return None


# ─────────────────────────── Entry / Gift Verification ───────────────────────

@login_required
def entry_page(request):
    config = _get_config()
    result = None
    query_id = ''

    if request.method == 'POST':
        query_id = request.POST.get('registration_id', '').strip()
        if query_id:
            child = _find_child_by_query(query_id)
            if child:
                day = config.current_day
                attendance = DailyAttendance.objects.filter(child=child, event_day=day).first()
                gift = DailyGift.objects.filter(child=child, event_day=day).first()
                result = {
                    'found': True,
                    'child': child,
                    'attendance': attendance,
                    'gift': gift,
                    'day': day,
                }
            else:
                result = {'found': False, 'query': query_id}

    return render(request, 'events/entry.html', {
        'config': config,
        'result': result,
        'query_id': query_id,
    })


@login_required
@require_POST
def record_entry(request, pk):
    child = get_object_or_404(Child, pk=pk)
    config = _get_config()
    day = config.current_day

    try:
        with transaction.atomic():
            attendance = DailyAttendance.objects.create(
                child=child,
                event_day=day,
                entry_time=timezone.now(),
            )
        return JsonResponse({
            'success': True,
            'message': f'✅ Entry recorded for {child.name}',
            'entry_time': attendance.entry_time.strftime('%I:%M %p'),
            'registration_id': child.registration_id,
        })
    except IntegrityError:
        existing = DailyAttendance.objects.filter(child=child, event_day=day).first()
        entry_time = existing.entry_time.strftime('%I:%M %p') if existing else 'Unknown'
        return JsonResponse({
            'success': False,
            'error': 'DUPLICATE',
            'message': f'⚠️ {child.name} already entered today at {entry_time}',
            'entry_time': entry_time,
        }, status=409)


@login_required
@require_POST
def record_gift(request, pk):
    child = get_object_or_404(Child, pk=pk)
    config = _get_config()
    day = config.current_day

    try:
        with transaction.atomic():
            gift = DailyGift.objects.create(
                child=child,
                event_day=day,
                gift_time=timezone.now(),
                status='GIVEN',
                operator=request.user,
            )
        return JsonResponse({
            'success': True,
            'message': f'🎁 Gift given to {child.name}',
            'gift_time': gift.gift_time.strftime('%I:%M %p'),
            'registration_id': child.registration_id,
        })
    except IntegrityError:
        existing = DailyGift.objects.filter(child=child, event_day=day).first()
        gift_time = existing.gift_time.strftime('%I:%M %p') if existing else 'Unknown'
        return JsonResponse({
            'success': False,
            'error': 'DUPLICATE',
            'message': f'⚠️ Gift already given to {child.name} today at {gift_time}',
            'gift_time': gift_time,
        }, status=409)


# API for live search (AJAX)
@login_required
def api_lookup(request):
    reg_id = request.GET.get('id', '').strip()
    if not reg_id:
        return JsonResponse({'found': False, 'error': 'No ID provided'})

    config = _get_config()
    day = config.current_day

    child = _find_child_by_query(reg_id)
    if child:
        attendance = DailyAttendance.objects.filter(child=child, event_day=day).first()
        gift = DailyGift.objects.filter(child=child, event_day=day).first()

        return JsonResponse({
            'found': True,
            'child': {
                'pk': child.pk,
                'registration_id': child.registration_id,
                'name': child.name,
                'guardian_name': child.guardian_name,
                'phone': child.phone,
                'photo_url': child.photo.url if child.photo else '',
                'age': child.age,
                'gender': child.get_gender_display() if child.gender else '',
            },
            'day': day,
            'attendance': {
                'recorded': bool(attendance),
                'entry_time': attendance.entry_time.strftime('%I:%M %p') if attendance else None,
            },
            'gift': {
                'given': bool(gift),
                'gift_time': gift.gift_time.strftime('%I:%M %p') if gift else None,
            },
        })
    else:
        return JsonResponse({'found': False, 'error': f'No child with ID {reg_id}'})


# ─────────────────────────── Search ──────────────────────────────────────────

@login_required
def search(request):
    query = request.GET.get('q', '').strip()
    results = []

    if query:
        # Search by registration ID (supports plain '1', '2026-1', partial), name, guardian, or phone
        id_q = Q(registration_id__icontains=query)
        if query.isdigit():
            id_q |= Q(registration_id__iexact=f"2026-{query}") | Q(registration_id__icontains=f"2026-{query}")
        elif query.startswith("2026-"):
            id_q |= Q(registration_id__iexact=query[5:])

        children = Child.objects.filter(
            id_q |
            Q(name__icontains=query) |
            Q(guardian_name__icontains=query) |
            Q(phone__icontains=query)
        ).order_by('id')
        results = list(children)

    return render(request, 'events/search.html', {
        'query': query,
        'results': results,
    })


# ─────────────────────────── Reports ─────────────────────────────────────────

@login_required
def daily_report(request):
    config = _get_config()
    selected_day = int(request.GET.get('day', config.current_day))
    selected_day = max(1, min(selected_day, config.total_days))

    attendances = {a.child_id: a for a in DailyAttendance.objects.filter(event_day=selected_day).select_related('child')}
    gifts = {g.child_id: g for g in DailyGift.objects.filter(event_day=selected_day).select_related('child')}

    all_children = Child.objects.all()
    report_rows = []
    for child in all_children:
        report_rows.append({
            'child': child,
            'attendance': attendances.get(child.pk),
            'gift': gifts.get(child.pk),
        })

    return render(request, 'events/daily_report.html', {
        'config': config,
        'selected_day': selected_day,
        'report_rows': report_rows,
        'total_days': range(1, config.total_days + 1),
        'entries_count': len(attendances),
        'gifts_count': len(gifts),
    })


# ─────────────────────────── Event Config ────────────────────────────────────

@login_required
def event_settings(request):
    config = _get_config()
    if request.method == 'POST':
        form = EventConfigForm(request.POST, instance=config)
        if form.is_valid():
            form.save()
            messages.success(request, 'Event settings updated successfully.')
            return redirect('event_settings')
    else:
        form = EventConfigForm(instance=config)

    return render(request, 'events/event_settings.html', {
        'form': form,
        'config': config,
        'day_choices': range(1, config.total_days + 1),
    })


@login_required
@require_POST
def change_day(request):
    """Quick day change from navbar."""
    day = int(request.POST.get('day', 1))
    config = _get_config()
    day = max(1, min(day, config.total_days))
    config.current_day = day
    config.save()
    messages.success(request, f'Current day changed to Day {day}.')
    return redirect(request.META.get('HTTP_REFERER', 'dashboard'))


# ─────────────────────────── Export ──────────────────────────────────────────

@login_required
def export_data(request):
    export_type = request.GET.get('type', '')
    config = _get_config()

    if not export_type:
        return render(request, 'events/export.html', {'config': config})

    if export_type == 'children':
        return _export_children()
    elif export_type == 'today_attendance':
        return _export_today_attendance(config.current_day)
    elif export_type == 'today_gifts':
        return _export_today_gifts(config.current_day)
    elif export_type == 'all_attendance':
        return _export_all_attendance()
    elif export_type == 'all_gifts':
        return _export_all_gifts()
    else:
        return HttpResponse('Invalid export type', status=400)


def _export_children():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Registered Children'

    headers = ['Registration ID', 'Child Name', 'Guardian Name', 'Phone', 'Age', 'Gender', 'Notes', 'Registered At']
    _style_headers(ws, headers)

    for child in Child.objects.all().order_by('registration_id'):
        ws.append([
            child.registration_id,
            child.name,
            child.guardian_name,
            child.phone,
            child.age or '',
            child.get_gender_display() if child.gender else '',
            child.notes,
            child.created_at.strftime('%Y-%m-%d %H:%M'),
        ])

    return _xlsx_response(wb, 'children_registered')


def _export_today_attendance(day):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f'Day {day} Attendance'

    headers = ['Registration ID', 'Child Name', 'Guardian Name', 'Phone', 'Entry Time']
    _style_headers(ws, headers)

    for a in DailyAttendance.objects.filter(event_day=day).select_related('child').order_by('entry_time'):
        ws.append([
            a.child.registration_id,
            a.child.name,
            a.child.guardian_name,
            a.child.phone,
            a.entry_time.strftime('%Y-%m-%d %H:%M:%S'),
        ])

    return _xlsx_response(wb, f'day{day}_attendance')


def _export_today_gifts(day):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f'Day {day} Gifts'

    headers = ['Registration ID', 'Child Name', 'Guardian Name', 'Phone', 'Gift Time', 'Status']
    _style_headers(ws, headers)

    for g in DailyGift.objects.filter(event_day=day).select_related('child').order_by('gift_time'):
        ws.append([
            g.child.registration_id,
            g.child.name,
            g.child.guardian_name,
            g.child.phone,
            g.gift_time.strftime('%Y-%m-%d %H:%M:%S'),
            g.status,
        ])

    return _xlsx_response(wb, f'day{day}_gifts')


def _export_all_attendance():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'All Attendance'

    headers = ['Registration ID', 'Child Name', 'Guardian Name', 'Phone', 'Event Day', 'Entry Time']
    _style_headers(ws, headers)

    for a in DailyAttendance.objects.select_related('child').order_by('event_day', 'entry_time'):
        ws.append([
            a.child.registration_id,
            a.child.name,
            a.child.guardian_name,
            a.child.phone,
            a.event_day,
            a.entry_time.strftime('%Y-%m-%d %H:%M:%S'),
        ])

    return _xlsx_response(wb, 'all_attendance_history')


def _export_all_gifts():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'All Gifts'

    headers = ['Registration ID', 'Child Name', 'Guardian Name', 'Phone', 'Event Day', 'Gift Time', 'Status']
    _style_headers(ws, headers)

    for g in DailyGift.objects.select_related('child').order_by('event_day', 'gift_time'):
        ws.append([
            g.child.registration_id,
            g.child.name,
            g.child.guardian_name,
            g.child.phone,
            g.event_day,
            g.gift_time.strftime('%Y-%m-%d %H:%M:%S'),
            g.status,
        ])

    return _xlsx_response(wb, 'all_gifts_history')


def _style_headers(ws, headers):
    """Apply orange header style to worksheet."""
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill(start_color='D4430A', end_color='D4430A', fill_type='solid')
        cell.alignment = Alignment(horizontal='center')
    ws.freeze_panes = 'A2'


def _xlsx_response(wb, filename):
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
    wb.save(response)
    return response


# ─────────────────────────── Home redirect ───────────────────────────────────

def home(request):
    if request.user.is_authenticated:
        return redirect('dashboard')
    return redirect('login')
